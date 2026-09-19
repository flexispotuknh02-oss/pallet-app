import os
import re
import sys
import time
from threading import Thread
from datetime import datetime
from traceback import print_exc
import tkinter as tk
from tkinter import messagebox
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

VER = '2.7.0'

XLSM_NAME = '自动化上架计划 (已修复).xlsm'
PLAN_SHEET = '2-上架计划'


def get_now_date_str():
    # 格式：2026.09.19
    return time.strftime('%Y.%m.%d', time.localtime())


def get_desktop_path():
    return os.path.join(os.path.expanduser("~"), "Desktop")


def data_format(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime('%Y/%m/%d')
    return str(value).strip()


# ---------------------------------------------------------------
# 柜号读取
# ---------------------------------------------------------------
def clean_container_no(value):
    """去掉文件名里不允许的字符和空白；空值返回 None"""
    if value is None:
        return None
    text = re.sub(r'[\\/:*?"<>|\s]', '', str(value))
    return text or None


# 标准集装箱号（ISO 6346）：3 位箱主代码 + 1 位设备类别(U/J/Z/R) + 6 位序列号 + 1 位校验位
CONTAINER_RE = re.compile(r'(?<![A-Z0-9])[A-Z]{3}[UJZR]\d{7}(?![A-Z0-9])')


def iso6346_check_ok(code):
    """校验位检查，用来排除偶然长得像柜号的其他编号"""
    values = {}
    v = 10
    for ch in 'ABCDEFGHIJKLMNOPQRSTUVWXYZ':
        if v % 11 == 0:      # 11、22、33 被跳过
            v += 1
        values[ch] = v
        v += 1
    total = 0
    for i, ch in enumerate(code[:10]):
        n = values[ch] if ch.isalpha() else int(ch)
        total += n * (2 ** i)
    return total % 11 % 10 == int(code[10])


def find_table_start(lines):
    """托盘明细表开始的位置（序号 1 后面紧跟托盘号）；找不到则返回 len(lines)"""
    n = len(lines)
    for i in range(n):
        if lines[i] == '1' and i + 1 < n and '-' in lines[i + 1]:
            return i
    return n


def find_container_no(lines):
    """
    从『收货详情.txt』表头里找柜号，与界面语言无关：
      ① 优先按标准集装箱号格式（4 个字母 + 7 位数字）在表头里查找，并用校验位排除误判；
      ② 找不到（比如运单号不是标准格式）时，取入库单号（ASN-开头的单独一行）往后第 2 行的内容。
    返回 (柜号或 None, 日志信息)
    """
    header = [ln.upper() for ln in lines[:find_table_start(lines)]]

    candidates = []
    for ln in header:
        for m in CONTAINER_RE.findall(ln):
            if m not in candidates:
                candidates.append(m)

    valid = [c for c in candidates if iso6346_check_ok(c)]
    if valid:
        extra = f"（另有其他候选 {[c for c in candidates if c != valid[0]]}，已取第一个）" if len(candidates) > 1 else ''
        return valid[0], f"从收货详情读取到柜号: {valid[0]}{extra}"
    if candidates:
        return candidates[0], f"从收货详情读取到柜号: {candidates[0]}（校验位不符，请核对是否正确）"

    # 兜底：ASN 单号所在行 + 2 行 = 集装箱号/运单号
    for i, ln in enumerate(header):
        if re.fullmatch(r'ASN-[A-Z0-9-]+', ln) and i + 2 < len(header):
            value = clean_container_no(header[i + 2])
            if value and re.fullmatch(r'[A-Z0-9-]{6,}', value):
                return value, f"从收货详情读取到单号: {value}（不是标准集装箱号格式，请核对）"

    return None, "收货详情里没有找到柜号"


def get_container_no_from_disk():
    """备用：从桌面已保存的 xlsm 文件读取 F1（只是上次保存时的内容，可能不是最新）"""
    xlsm_path = os.path.join(get_desktop_path(), XLSM_NAME)
    if not os.path.exists(xlsm_path):
        return None, f"未在桌面找到 '{XLSM_NAME}'"

    try:
        wb = load_workbook(xlsm_path, data_only=True)
        ws = wb[PLAN_SHEET] if PLAN_SHEET in wb.sheetnames else wb.active
        value = ws['F1'].value
        wb.close()
    except Exception as e:
        return None, f"读取磁盘文件失败: {e}"

    container_no = clean_container_no(value)
    if container_no is None:
        return None, "磁盘文件中 F1 单元格为空"
    return container_no, f"读取到磁盘文件里保存的柜号: {container_no}"


# ---------------------------------------------------------------
# 收货详情.txt
# ---------------------------------------------------------------
def read_receipt_lines():
    """读取桌面『收货详情.txt』，返回 (行列表或 None, 错误信息)"""
    txt_path = os.path.join(get_desktop_path(), '收货详情.txt')

    if not os.path.exists(txt_path):
        return None, f"未在桌面找到 '收货详情.txt' 文件，请检查路径: {txt_path}"

    try:
        with open(txt_path, 'r', encoding='utf-8') as f:
            lines = [line.strip() for line in f.readlines()]
    except UnicodeDecodeError:
        with open(txt_path, 'r', encoding='gbk') as f:
            lines = [line.strip() for line in f.readlines()]
    return lines, ''


def parse_receipt_pallets(lines):
    """解析托盘明细，返回 (托盘列表, 日志信息)"""
    system_pallets = []
    n = len(lines)
    i = find_table_start(lines)

    while i + 5 < n:
        seq = lines[i]
        pallet_full = lines[i+1]
        sku = lines[i+2]
        count_str = lines[i+3]

        if pallet_full.startswith('IB') and count_str.isdigit():
            last_3 = pallet_full.split('-')[-1].zfill(3)
            system_pallets.append({
                'pallet_last3': last_3,
                'pallet_full': pallet_full,
                'sku': sku,
                'count': int(count_str),
                'used': False
            })
            i += 6
        else:
            i += 1

    return system_pallets, f"成功解析桌面 '收货详情.txt'，共获取 {len(system_pallets)} 托数据。"


def build_pallet_mapping(excel_rows, system_pallets):
    excel_by_sku = {}
    for idx, row in enumerate(excel_rows):
        loc, sku, count = row[0], row[1], row[2]
        count_num = int(count) if str(count).isdigit() else 0

        if sku not in excel_by_sku:
            excel_by_sku[sku] = []
        excel_by_sku[sku].append({
            'excel_index': idx,
            'location': loc,
            'count': count_num
        })

    sys_by_sku = {}
    if system_pallets:
        for p in system_pallets:
            sku = p['sku']
            if sku not in sys_by_sku:
                sys_by_sku[sku] = []
            sys_by_sku[sku].append(p)

    mapping_result = {}

    for sku, ex_items in excel_by_sku.items():
        sys_items = sys_by_sku.get(sku, [])

        if not sys_items or len(ex_items) != len(sys_items):
            for item in ex_items:
                mapping_result[item['excel_index']] = "N/A"
            continue

        ex_counts = sorted([x['count'] for x in ex_items])
        sys_counts = sorted([x['count'] for x in sys_items])

        if ex_counts == sys_counts:
            temp_sys = list(sys_items)
            for ex in ex_items:
                matched = False
                for s in temp_sys:
                    if not s['used'] and s['count'] == ex['count']:
                        s['used'] = True
                        mapping_result[ex['excel_index']] = s['pallet_last3']
                        matched = True
                        break
                if not matched:
                    mapping_result[ex['excel_index']] = "N/A"
        else:
            def get_9th_char(item):
                loc = item['location']
                return loc[8] if len(loc) > 8 else loc

            sorted_ex = sorted(ex_items, key=get_9th_char)
            sorted_sys = sorted(sys_items, key=lambda x: x['count'])

            for ex, sys_p in zip(sorted_ex, sorted_sys):
                sys_p['used'] = True
                mapping_result[ex['excel_index']] = sys_p['pallet_last3']

    return mapping_result


def process_single(file_path, result_path, system_pallets):
    file_dir, _ = os.path.split(file_path)
    font_path = os.path.join(file_dir, 'calibrib.ttf')
    img_path = os.path.join(file_dir, 'company.png')

    if not os.path.isfile(font_path) or not os.path.isfile(img_path):
        raise Exception('缺少 calibrib.ttf 或 company.png 文件！')

    pdfmetrics.registerFont(TTFont('Calibri', font_path))
    my_canvas = canvas.Canvas(result_path, pagesize=(841.89, 595.27))

    table_w = 779.08
    table_h = 458.36
    table_x = 32.9
    table_y = 101.2

    lines_cord = [440.0, 320.0, 200.0, 150.0]
    ver_line_x = 340.0

    fields = ['Position:', 'SKU:', 'Quantity:', 'Inbound Order:', 'Date:']

    wb = load_workbook(file_path)
    wst = wb.active

    raw_excel_rows = []
    for row in wst.iter_rows(min_row=2):
        rec = [data_format(x) for x in row]
        if rec and rec[0]:
            raw_excel_rows.append(rec)

    pallet_mapping = build_pallet_mapping(raw_excel_rows, system_pallets)

    # 构建待生成页面列表
    pages_data = []
    for idx, record in enumerate(raw_excel_rows):
        location = record[0]
        sku = record[1] if len(record) > 1 else ""
        count = record[2] if len(record) > 2 else ""
        inbound_order = record[3] if len(record) > 3 else ""
        date_str = record[4] if len(record) > 4 else ""

        pallet_no = pallet_mapping.get(idx, "N/A")

        if pallet_no != "N/A":
            qty_pallet_display = f"{count}  (Pallet: #{pallet_no})"
        else:
            qty_pallet_display = f"{count}"

        pages_data.append({
            'pallet_no': pallet_no,
            'data_list': [location, sku, qty_pallet_display, inbound_order, date_str]
        })

    # 按托盘尾号数字排序，N/A 统一排到最后
    def get_sort_key(item):
        p = item['pallet_no']
        if p != "N/A" and p.isdigit():
            return (0, int(p))  # 0 保证数字排在前面，按整数大小升序
        return (1, 999999)      # 1 保证 N/A 排在最后

    pages_data.sort(key=get_sort_key)

    # 按照排序后的顺序绘制 PDF 页面
    for page in pages_data:
        data_list = page['data_list']

        # 1. 绘制外框
        my_canvas.rect(table_x, table_y, table_w, table_h)

        # 2. 绘制横线
        for line_y in lines_cord:
            my_canvas.line(table_x, line_y, table_x + table_w, line_y)

        # 3. 绘制垂直分割线
        my_canvas.line(ver_line_x, table_y, ver_line_x, table_y + table_h)

        # 4. 写入字段名与值
        y_positions = [465.0, 350.0, 230.0, 162.0, 113.0]

        for i in range(5):
            y = y_positions[i]
            val = data_list[i]

            # 左侧标题
            if i in [0, 1, 2]:
                my_canvas.setFont('Calibri', 52)
            else:
                my_canvas.setFont('Calibri', 24)
            my_canvas.drawString(table_x + 15, y, fields[i])

            # 右侧数据值
            if i in [0, 1]:  # Position, SKU
                if len(val) > 22:
                    my_canvas.setFont('Calibri', 28)
                elif len(val) > 14:
                    my_canvas.setFont('Calibri', 38)
                else:
                    my_canvas.setFont('Calibri', 48)
                my_canvas.drawString(ver_line_x + 15, y, val)

            elif i == 2:  # Quantity + Pallet
                if len(val) > 15:
                    my_canvas.setFont('Calibri', 36)
                else:
                    my_canvas.setFont('Calibri', 52)
                my_canvas.drawString(ver_line_x + 15, y, val)

            else:  # Inbound Order, Date
                my_canvas.setFont('Calibri', 26)
                my_canvas.drawString(ver_line_x + 15, y, val)

        # 5. 画底部 Logo
        my_canvas.drawInlineImage(img_path, 230.29, 18.0, width=307.598, height=75.8965)
        my_canvas.showPage()

    my_canvas.save()


def thread_exc(func, *args, **kwargs):
    t = Thread(target=func, args=args, kwargs=kwargs)
    t.daemon = True
    t.start()


def process(file_dir, log_widget):
    lines, read_err = read_receipt_lines()

    if lines is None:
        log_widget.insert(tk.END, f'{read_err}\n')
        log_widget.insert(tk.END, '无法读取桌面 txt，将默认以 N/A 生成标签！\n')
        system_pallets = []
        container_no = None
    else:
        system_pallets, txt_msg = parse_receipt_pallets(lines)
        log_widget.insert(tk.END, f'{txt_msg}\n')
        container_no, container_msg = find_container_no(lines)
        log_widget.insert(tk.END, f'{container_msg}\n')

    # 收货详情里没找到柜号时，才退回读取磁盘上的 xlsm（可能是旧的，需要用户确认）
    if container_no is None:
        disk_no, disk_msg = get_container_no_from_disk()
        log_widget.insert(tk.END, f'{disk_msg}\n')
        if disk_no:
            log_widget.insert(tk.END, '注意：磁盘文件里的柜号可能不是最新的！\n')
            use_it = messagebox.askyesno(
                '确认柜号',
                f'收货详情里没有找到柜号。\n\n磁盘上已保存的 Excel 里的柜号是：{disk_no}\n'
                f'如果你还没保存过最新的柜号，这个号码可能是旧的。\n\n'
                f'是否仍用它命名 PDF？\n（选"否"则改用原 Excel 文件名命名）'
            )
            if use_it:
                container_no = disk_no

    if container_no is None:
        log_widget.insert(tk.END, '将使用原 Excel 文件名命名 PDF。\n')

    result_dir = os.path.join(file_dir, 'pallet pdf')
    if not os.path.isdir(result_dir):
        os.makedirs(result_dir)

    log_widget.insert(tk.END, '开始生成 PDF...\n')
    top.update()
    start = time.time()
    used_names = set()  # 防止同一次运行里多个 xlsx 生成同名 PDF 互相覆盖

    for file_name in os.listdir(file_dir):
        if file_name.endswith('.xlsx') and not file_name.startswith('~$'):
            file_path = os.path.join(file_dir, file_name)

            if container_no:
                base_name = f'{get_now_date_str()} - {container_no}'
            else:
                base_name = file_name.replace('.xlsx', '') + f'_{get_now_date_str()}'

            final_name, n = base_name, 2
            while final_name in used_names:
                final_name = f'{base_name}_{n}'
                n += 1
            used_names.add(final_name)
            result_path = os.path.join(result_dir, final_name + '.pdf')

            try:
                process_single(file_path, result_path, system_pallets)
                info = f'{file_name} 处理完成 -> {final_name}.pdf\n'
            except Exception as e:
                print_exc()
                info = f'{file_name} 处理失败: {str(e)}\n'

            log_widget.insert(tk.END, info)
            top.update()

    spend_time = round(time.time() - start, 1)
    log_widget.insert(tk.END, f'全部完成，耗时 {spend_time} 秒！\n')
    top.update()


top = tk.Tk()
top.title(f'Loctek 托盘标签自动生成工具 v{VER}')
top.geometry('550x360')


def main_init():
    tk.Label(top, text='Excel 文件夹路径:').place(x=20, y=25)
    dir_input = tk.Entry(top, font=('Arial', 10), width=42)
    dir_input.insert(0, './')
    dir_input.place(x=140, y=25)

    tk.Label(top, text='提示: 运行前请确保已将最新文本保存到桌面 [收货详情.txt]', fg='gray').place(x=20, y=60)

    tk.Label(top, text='运行日志:').place(x=20, y=95)
    log_box = tk.Text(top, font=('Arial', 9), height=11, width=70)
    log_box.place(x=20, y=120)

    def check_and_run():
        dir_path = dir_input.get().strip()
        if not os.path.isdir(dir_path):
            messagebox.showinfo('提示', '请输入有效的文件夹路径！')
            return
        process(dir_path, log_box)

    btn = tk.Button(top, text='开始生成标签', bg='#4CAF50', fg='white', font=('Arial', 11, 'bold'), command=lambda: thread_exc(check_and_run))
    btn.place(x=190, y=295, width=170, height=38)


if __name__ == '__main__':
    main_init()
    top.mainloop()
